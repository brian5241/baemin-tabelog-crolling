"""
배민 업체 수집기 v2.2
BlueStacks의 배민 앱에서 업체 목록을 자동으로 수집합니다.

사용법:
  py scraper.py                                수집 시작
  py scraper.py --test                         화면 한 번만 캡처해서 결과 확인
  py scraper.py --count 50                     50개 업체 수집
  py scraper.py --categories 양식 중식 한식      여러 카테고리 연속 수집
  py scraper.py --schedule 14:30               매일 14:30에 자동 수집
  py scraper.py --repeat 90                    90분 간격 반복 수집
  py scraper.py --repeat 60 --categories 양식 중식  60분 간격 + 멀티카테고리

config.json의 ocr.engine으로 엔진 선택:
  "easyocr"    — 로컬 OCR (기본값, GPU 없어도 동작)
  "anthropic"  — Claude API (빠르고 정확, api_key 필요)
"""

import sys
import os
import re
import time
import json
import random
import base64
import io
import ctypes
import ctypes.wintypes
from datetime import datetime
from pathlib import Path

ctypes.windll.user32.SetProcessDPIAware()
_hdc = ctypes.windll.user32.GetDC(0)
_dpi = ctypes.windll.gdi32.GetDeviceCaps(_hdc, 88)
ctypes.windll.user32.ReleaseDC(0, _hdc)
DPI_SCALE = _dpi / 96.0

SCRIPT_DIR = Path(__file__).parent
OUTPUT_DIR = SCRIPT_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)
CONFIG_PATH = SCRIPT_DIR / "config.json"
CORRECTIONS_PATH = SCRIPT_DIR / "corrections.json"

# --------------- OCR 후처리 사전 ---------------

_DEFAULT_CORRECTIONS = {
    "배스권라반스": "배스킨라빈스",
    "배스킨라빈스": "배스킨라빈스",
    "배스권라빈스": "배스킨라빈스",
    "스타버스": "스타벅스",
    "스타벅스": "스타벅스",
    "스타벅츠": "스타벅스",
    "도미노피자": "도미노피자",
    "맥도넛드": "맥도날드",
    "맥도날드": "맥도날드",
    "맥도냘드": "맥도날드",
    "버거킹": "버거킹",
    "버거킹그": "버거킹",
    "롯데리야": "롯데리아",
    "롯데리아": "롯데리아",
    "뚜레주르": "뚜레쥬르",
    "뚜레쥬르": "뚜레쥬르",
    "파리바겟뜨": "파리바게뜨",
    "파리바게뜨": "파리바게뜨",
    "파리바게트": "파리바게뜨",
    "교촌치킨": "교촌치킨",
    "교촌치견": "교촌치킨",
    "비비큐": "BBQ",
    "비비규": "BBQ",
    "네네치견": "네네치킨",
    "네네치킨": "네네치킨",
    "굽네치견": "굽네치킨",
    "굽네치킨": "굽네치킨",
    "신전먹볶이": "신전떡볶이",
    "신전떡볶이": "신전떡볶이",
    "뿌떡먹볶이": "뿌떡떡볶이",
    "본죽&비범밥": "본죽&비빔밥",
    "본죽&비빔밥": "본죽&비빔밥",
    "하삼동커피": "하삼동커피",
    "공차": "공차",
    "이디야커피": "이디야커피",
    "이디야카피": "이디야커피",
    "메가커피": "메가커피",
    "메가카피": "메가커피",
    "컴포즈커피": "컴포즈커피",
    "컴포즈카피": "컴포즈커피",
    "피자헛": "피자헛",
    "피자헏": "피자헛",
    "한솥도시럭": "한솥도시락",
    "한솥도시락": "한솥도시락",
    "놀부": "놀부",
    "김밥천국": "김밥천국",
    "김밥천귝": "김밥천국",
    "세수대야방수": "세수대야빙수",
    "왕방수": "왕빙수",
    "일산단현점": "일산탄현점",
    "단현점": "탄현점",
    "문존점": "문촌점",
    "주업점": "주엽점",
    "후곡마을점": "후곡마을점",
    "해월카메": "해월카페",
    "흙 메이드": "홈 메이드",
    "피자스굴": "피자스쿨",
    "피자첫": "피자헛",
    "던권": "던킨",
    "던견": "던킨",
    "Gh": "GH",
}

def _load_corrections():
    corrections = dict(_DEFAULT_CORRECTIONS)
    if CORRECTIONS_PATH.exists():
        try:
            with open(str(CORRECTIONS_PATH), encoding='utf-8') as f:
                user_corrections = json.load(f)
            corrections.update(user_corrections)
        except Exception:
            pass
    return corrections

OCR_CORRECTIONS = _load_corrections()


def load_config():
    if CONFIG_PATH.exists():
        with open(str(CONFIG_PATH), encoding='utf-8') as f:
            return json.load(f)
    return {}


CFG = load_config()


def cfg(section, key, default=None):
    return CFG.get(section, {}).get(key, default)


# --------------- OCR 엔진 ---------------

def _init_easyocr():
    import easyocr
    langs = cfg('ocr', 'languages', ['ko', 'en'])
    gpu = cfg('ocr', 'gpu', False)
    return easyocr.Reader(langs, gpu=gpu)


def _read_easyocr(reader, image):
    import numpy as np
    w, h = image.size
    min_width = cfg('ocr', 'min_width', 540)
    if w < min_width:
        scale = min_width / w
        image = image.resize((int(w * scale), int(h * scale)), resample=3)
        print(f"  이미지 확대: {w}x{h} -> {image.size[0]}x{image.size[1]}")

    arr = np.array(image)
    mag = cfg('ocr', 'mag_ratio', 1.5)
    retry_mags = cfg('ocr', 'retry_mag_ratios', [2.0, 2.5])
    min_texts = cfg('ocr', 'min_texts_before_retry', 5)

    results = reader.readtext(arr, mag_ratio=mag)
    if len(results) >= min_texts:
        return [(bbox, text, conf) for bbox, text, conf in results]

    for alt_mag in retry_mags:
        time.sleep(random.uniform(0.3, 0.8))
        alt_results = reader.readtext(arr, mag_ratio=alt_mag)
        if len(alt_results) > len(results):
            print(f"  mag_ratio {mag} -> {alt_mag} 재시도: {len(results)} -> {len(alt_results)}개 텍스트")
            results = alt_results
            if len(results) >= min_texts:
                break

    return [(bbox, text, conf) for bbox, text, conf in results]


def _init_anthropic():
    try:
        import anthropic
    except ImportError:
        print("anthropic 패키지가 없습니다: py -m pip install anthropic")
        sys.exit(1)
    key = cfg('anthropic', 'api_key', '')
    if not key:
        print("config.json의 anthropic.api_key를 설정하세요.")
        sys.exit(1)
    return anthropic.Anthropic(api_key=key)


def _read_anthropic(client, image):
    model = cfg('anthropic', 'model', 'claude-sonnet-4-20250514')
    buf = io.BytesIO()
    image.save(buf, format='PNG')
    b64 = base64.standard_b64encode(buf.getvalue()).decode()

    resp = client.messages.create(
        model=model,
        max_tokens=2048,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {
                    "type": "base64", "media_type": "image/png", "data": b64}},
                {"type": "text", "text": (
                    "이 배민 앱 스크린샷에서 음식점 목록을 읽어주세요. "
                    "각 음식점의 이름, 평점(별점), 리뷰 수를 JSON 배열로 반환하세요. "
                    "형식: [{\"name\": \"...\", \"rating\": 4.9, \"reviews\": 786}, ...] "
                    "음식점이 아닌 UI 요소(버튼, 배너, 탭)는 제외하세요. "
                    "JSON만 반환하고 다른 텍스트는 쓰지 마세요."
                )},
            ],
        }],
    )

    text = resp.content[0].text.strip()
    if text.startswith('```'):
        text = re.sub(r'^```\w*\n?', '', text)
        text = re.sub(r'\n?```$', '', text)

    businesses = json.loads(text)
    return [_biz_to_ocr(b) for b in businesses]


def _biz_to_ocr(biz):
    name = biz.get('name', '')
    rating = biz.get('rating', '')
    reviews = biz.get('reviews', '')
    fake_text = f"{name} {rating} ({reviews})"
    fake_bbox = [[0, 0], [100, 0], [100, 20], [0, 20]]
    return (fake_bbox, fake_text, 0.99)


def _init_tesseract():
    try:
        import pytesseract
    except ImportError:
        print("pytesseract 패키지가 없습니다: py -m pip install pytesseract")
        sys.exit(1)
    tess_path = cfg('tesseract', 'path', r'C:\Program Files\Tesseract-OCR\tesseract.exe')
    if not Path(tess_path).exists():
        print(f"Tesseract를 찾을 수 없습니다: {tess_path}")
        print("winget install UB-Mannheim.TesseractOCR 로 설치하세요.")
        sys.exit(1)
    pytesseract.pytesseract.tesseract_cmd = tess_path
    return pytesseract


def _read_tesseract(pytesseract, image):
    import numpy as np
    arr = np.array(image)
    lang = cfg('tesseract', 'lang', 'kor+eng')
    psm = cfg('tesseract', 'psm', 6)
    custom_config = f'--oem 1 --psm {psm}'

    data = pytesseract.image_to_data(arr, lang=lang, config=custom_config,
                                      output_type=pytesseract.Output.DICT)

    results = []
    n = len(data['text'])
    for i in range(n):
        text = data['text'][i].strip()
        conf = int(data['conf'][i])
        if not text or conf < 20:
            continue
        x = data['left'][i]
        y = data['top'][i]
        w = data['width'][i]
        h = data['height'][i]
        bbox = [[x, y], [x + w, y], [x + w, y + h], [x, y + h]]
        results.append((bbox, text, conf / 100.0))

    return results


ENGINES = {
    'easyocr': (_init_easyocr, _read_easyocr),
    'anthropic': (_init_anthropic, _read_anthropic),
    'tesseract': (_init_tesseract, _read_tesseract),
}


def create_reader():
    engine = cfg('ocr', 'engine', 'easyocr')
    if engine not in ENGINES:
        print(f"지원하지 않는 엔진: {engine}")
        print(f"사용 가능: {', '.join(ENGINES)}")
        sys.exit(1)
    init_fn, _ = ENGINES[engine]
    print(f"엔진: {engine}")
    return engine, init_fn()


def read_image(engine_name, reader, image):
    _, read_fn = ENGINES[engine_name]
    return read_fn(reader, image)


# --------------- 파싱 ---------------

def _correct_ocr_name(name):
    """OCR 후처리: 브랜드명 부분 일치로 교정."""
    for wrong, right in OCR_CORRECTIONS.items():
        if wrong in name:
            name = name.replace(wrong, right)
    return name


def parse_businesses(ocr_results, img_height=0, return_positions=False):
    if not ocr_results:
        return ([], {}) if return_positions else []

    items = []
    for bbox, text, conf in ocr_results:
        y_mid = (bbox[0][1] + bbox[2][1]) / 2
        x_left = bbox[0][0]
        h = abs(bbox[2][1] - bbox[0][1])
        items.append({
            'text': text.strip(), 'y': y_mid, 'x': x_left,
            'h': h, 'conf': conf,
        })

    if img_height > 0:
        top_cut = img_height * 0.05
        bot_cut = img_height * 0.98
        items = [it for it in items if top_cut < it['y'] < bot_cut]
    if not items:
        return ([], {}) if return_positions else []

    items.sort(key=lambda x: (x['y'], x['x']))

    lines = []
    cur = [items[0]]
    for it in items[1:]:
        if abs(it['y'] - cur[0]['y']) <= 14:
            cur.append(it)
        else:
            lines.append(cur)
            cur = [it]
    if cur:
        lines.append(cur)

    joined = []
    line_ys = []
    line_confs = []
    for line in lines:
        line.sort(key=lambda x: x['x'])
        text = ' '.join(it['text'] for it in line)
        avg_conf = sum(it['conf'] for it in line) / len(line)
        joined.append(text)
        line_ys.append(line[0]['y'])
        line_confs.append(avg_conf)

    ui_skip = re.compile(
        r'기본순|할인.쿠폰|배달방식|별점|최소주문금액|기타$|'
        r'음식배달|배민클럽|배민|카페.디저트|패스트푸드|분식|중식$|한식$|'
        r'피자$|찜.탕|돈까스|홈$|장보기|쇼핑|주문내역|마이배민|'
        r'곧 사라|이번 주|한정|쿠폰 확인|배달팁|배달립|배달답|전체보기|'
        r'적용 중|배달 0원|0원 가게|^\d+:\d+|'
        r'^경기|^서울|^인천|^부산|^대구|^광주|^대전|'
        r'BlueStacks|App Player|'
        r'^\d[\d,.]+원\s*$|^[\d,.]+\s*$'
    )

    status_strip = re.compile(
        r'\s*(영업\s*준비\s*중|오늘\s*휴무|영업\s*종료|영업\s*시작\s*전|'
        r'일시\s*중지|주문\s*불가|준비\s*중)\s*'
    )

    biz_re = re.compile(
        r'(.+?)\s*[★☆\*]?\s*(\d[.,]\d)\s*[\(\[（]\s*([\d,.]+)\s*[\)\]）]'
    )
    biz_re_rev = re.compile(
        r'(\d[.,]\d)\s*[\(\[（]\s*([\d,.]+)\s*[\)\]）]\s+(.+)'
    )
    biz_re2 = re.compile(
        r'^(.{2,}?)\s+[\(\[（]\s*(\d{2,6})\s*[\)\]）]\s*$'
    )

    def _fix_rating_parens(t):
        t = re.sub(r'(\d\.\d)\s+(\d[\d,.]*)\)', r'\1 (\2)', t)
        t = re.sub(r'(\d\.\d\([\d,.]+?)[1lI|]\s*$', r'\1)', t)
        t = re.sub(r'(\d\.\d\([\d,.]+)\s*$', r'\1)', t)
        t = re.sub(r'(\d\.\d)\s+(\d[\d,.]{2,})\s*$', r'\1 (\2)', t)
        return t

    has_ko = re.compile(r'[가-힣]')

    def _clean_name(name):
        name = re.sub(r'\s*[\(\[（]\d[\d,.]*[\)\]）]\s*', ' ', name)
        name = re.sub(r'\s+\d\s+', ' ', name)
        name = re.sub(r'[★☆\*\s]+$', '', name).strip()
        name = re.sub(r'\s*광고\s*$', '', name).strip()
        name = re.sub(r'^\d+\s+', '', name)
        name = ' '.join(name.split())
        return _correct_ocr_name(name)

    promo_re = re.compile(
        r'놓치면.*할인|아까운.*할인|로컬\s*맛집|방금\s*본\s*가게|'
        r'비슷해요|골라담기|이\s*동네.*맛집|즐겨찾기|많이\s*주문한|'
        r'할인\s*가게|재주문|2천원|10%\s*이상|'
        r'곧\s*사라|평점.*이상.*가게|단골.*찜|이번\s*주.*한정'
    )

    skip_zones = []
    for idx, text in enumerate(joined):
        if promo_re.search(text):
            skip_zones.append(line_ys[idx])

    def _in_promo_zone(y):
        for zone_y in skip_zones:
            if zone_y + 25 < y <= zone_y + 250:
                return True
        return False

    businesses = []
    null_positions = {}
    low_conf_positions = {}
    for idx, text in enumerate(joined):
        if ui_skip.search(text):
            continue

        if _in_promo_zone(line_ys[idx]):
            continue

        text = status_strip.sub(' ', text).strip()
        text = _fix_rating_parens(text)

        found = False
        for m in biz_re.finditer(text):
            found = True
            name = _clean_name(m.group(1))
            rating = float(m.group(2).replace(',', '.'))
            reviews = int(m.group(3).replace(',', '').replace('.', ''))
            if not (1.0 <= rating <= 5.0) or reviews < 1:
                continue
            if len(name) < 2 or not has_ko.search(name):
                continue
            businesses.append({
                'name': name, 'rating': rating, 'reviews': reviews,
                'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M'),
            })
            if line_confs[idx] < 0.65:
                low_conf_positions[name] = line_ys[idx]

        if not found:
            m_rev = biz_re_rev.search(text)
            if m_rev:
                rating = float(m_rev.group(1).replace(',', '.'))
                reviews = int(m_rev.group(2).replace(',', '').replace('.', ''))
                name = m_rev.group(3)
                name = re.sub(r'\s+[가-힣]고\d*\s*$', '', name)
                name = _clean_name(name)
                if (1.0 <= rating <= 5.0 and reviews >= 1
                        and len(name) >= 2 and has_ko.search(name)):
                    businesses.append({
                        'name': name, 'rating': rating, 'reviews': reviews,
                        'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M'),
                    })
                    if line_confs[idx] < 0.65:
                        low_conf_positions[name] = line_ys[idx]
                    found = True

        if not found:
            m2 = biz_re2.match(text)
            if m2:
                name = _clean_name(m2.group(1))
                reviews = int(m2.group(2))
                if len(name) >= 2 and reviews >= 10 and has_ko.search(name) and not re.match(r'^[\d.,\s]+$', name):
                    businesses.append({
                        'name': name, 'rating': None, 'reviews': reviews,
                        'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M'),
                    })
                    null_positions[name] = line_ys[idx]

    if return_positions:
        return businesses, null_positions, low_conf_positions
    return businesses


# --------------- 화면 제어 ---------------

_bs_hwnd = None


def _bring_to_front():
    import ctypes
    global _bs_hwnd
    if _bs_hwnd:
        ctypes.windll.user32.SetForegroundWindow(_bs_hwnd)
        time.sleep(0.3)


def find_bluestacks():
    global _bs_hwnd

    hwnds = []
    def callback(hwnd, _):
        length = ctypes.windll.user32.GetWindowTextLengthW(hwnd)
        if length:
            buf = ctypes.create_unicode_buffer(length + 1)
            ctypes.windll.user32.GetWindowTextW(hwnd, buf, length + 1)
            if 'bluestacks app player' in buf.value.lower():
                hwnds.append((hwnd, buf.value))
        return True
    WNDENUMPROC = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)
    ctypes.windll.user32.EnumWindows(WNDENUMPROC(callback), 0)

    if not hwnds:
        return None

    hwnd, title = hwnds[0]
    _bs_hwnd = hwnd

    ctypes.windll.user32.ShowWindow(hwnd, 9)
    time.sleep(0.5)
    ctypes.windll.user32.SetForegroundWindow(hwnd)
    time.sleep(0.5)

    rect = ctypes.wintypes.RECT()
    ctypes.windll.user32.GetWindowRect(hwnd, ctypes.byref(rect))

    if rect.left < -10000:
        ctypes.windll.user32.ShowWindow(hwnd, 9)
        time.sleep(1.0)
        ctypes.windll.user32.GetWindowRect(hwnd, ctypes.byref(rect))

    title_bar = int(35 * DPI_SCALE)
    top = max(0, rect.top + title_bar)
    height = rect.bottom - top
    width = rect.right - rect.left

    if height < 100 or width < 100:
        print("오류: BlueStacks 창이 너무 작습니다.")
        return None

    return {
        'title': title,
        'left': rect.left,
        'top': top,
        'width': width,
        'height': height,
    }


def capture(region):
    import win32gui, win32ui
    from PIL import Image
    global _bs_hwnd

    if _bs_hwnd:
        rect = ctypes.wintypes.RECT()
        ctypes.windll.user32.GetWindowRect(_bs_hwnd, ctypes.byref(rect))
        w = rect.right - rect.left
        h = rect.bottom - rect.top
        if w > 100 and h > 100:
            wDC = win32gui.GetWindowDC(_bs_hwnd)
            dcObj = win32ui.CreateDCFromHandle(wDC)
            cDC = dcObj.CreateCompatibleDC()
            bmp = win32ui.CreateBitmap()
            bmp.CreateCompatibleBitmap(dcObj, w, h)
            cDC.SelectObject(bmp)
            ctypes.windll.user32.PrintWindow(_bs_hwnd, cDC.GetSafeHdc(), 2)
            bmpinfo = bmp.GetInfo()
            bmpstr = bmp.GetBitmapBits(True)
            img = Image.frombuffer('RGB',
                (bmpinfo['bmWidth'], bmpinfo['bmHeight']),
                bmpstr, 'raw', 'BGRX', 0, 1)
            dcObj.DeleteDC()
            cDC.DeleteDC()
            win32gui.ReleaseDC(_bs_hwnd, wDC)
            win32gui.DeleteObject(bmp.GetHandle())
            title_bar = int(35 * DPI_SCALE)
            return img.crop((0, title_bar, w, h))

    import pyautogui
    _bring_to_front()
    return pyautogui.screenshot(region=(
        region['left'], region['top'],
        region['width'], region['height'],
    ))


def human_pause():
    chance = cfg('scroll', 'long_pause_chance', 0.12)
    if random.random() < chance:
        lo = cfg('scroll', 'long_pause_min', 5.0)
        hi = cfg('scroll', 'long_pause_max', 9.0)
        wait = random.uniform(lo, hi)
        print(f"  (잠시 쉬는 중... {wait:.0f}초)")
        time.sleep(wait)
    else:
        lo = cfg('scroll', 'pause_min', 1.8)
        hi = cfg('scroll', 'pause_max', 3.5)
        time.sleep(random.uniform(lo, hi))


def _category_pause():
    """카테고리 전환 시 사람처럼 긴 쉼. 메뉴를 훑어보는 느낌."""
    wait = random.uniform(3.0, 6.0)
    print(f"  (카테고리 둘러보는 중... {wait:.0f}초)")
    time.sleep(wait)


KNOWN_TABS = ['전체', '치킨', '한식', '중식', '분식', '카페·디저트',
              '피자', '패스트푸드', '양식', '일식', '고기·구이', '야식',
              '찜·탕', '돈까스', '족발·보쌈', '도시락', '회·초밥']


def _scroll_up(region, count=3):
    """위로 부드럽게 스크롤 (pull-to-refresh 방지용 짧은 거리)."""
    import pyautogui
    _bring_to_front()
    cx = region['left'] + region['width'] // 2
    for _ in range(count):
        y_start = region['top'] + int(region['height'] * 0.40)
        y_end = region['top'] + int(region['height'] * 0.65)
        pyautogui.moveTo(cx + random.randint(-10, 10), y_start,
                         duration=random.uniform(0.08, 0.15))
        pyautogui.mouseDown()
        pyautogui.moveTo(cx + random.randint(-5, 5), y_end,
                         duration=random.uniform(0.2, 0.35))
        pyautogui.mouseUp()
        time.sleep(random.uniform(0.3, 0.5))
    time.sleep(random.uniform(0.3, 0.6))


def _swipe_tab_bar(region, direction='left'):
    import pyautogui
    _bring_to_front()
    tab_y = region['top'] + int(region['height'] * 0.17)
    if direction == 'left':
        x_start = region['left'] + region['width'] - 50
        x_end = region['left'] + 50
    else:
        x_start = region['left'] + 50
        x_end = region['left'] + region['width'] - 50
    pyautogui.moveTo(x_start, tab_y + random.randint(-3, 3),
                     duration=random.uniform(0.15, 0.25))
    time.sleep(random.uniform(0.05, 0.1))
    pyautogui.mouseDown()
    pyautogui.moveTo(x_end, tab_y + random.randint(-3, 3),
                     duration=random.uniform(0.3, 0.5))
    pyautogui.mouseUp()
    time.sleep(random.uniform(0.5, 0.8))


def _img_to_screen(region, img, ocr_x, ocr_y, ocr_scale):
    """OCR 좌표를 screen 좌표로 변환 (DPI 스케일 보정 포함)."""
    w, h = img.size
    dpi_x = w / region['width']
    dpi_y = h / region['height']
    return (
        region['left'] + int(ocr_x / ocr_scale / dpi_x),
        region['top'] + int(ocr_y / ocr_scale / dpi_y),
    )


def _find_category_tab(region, category_name, engine_name, reader):
    """탭 좌표를 찾으면 (pos, img, results), 못 찾으면 (None, img, results) 반환."""
    img = capture(region)
    results = read_image(engine_name, reader, img)

    w, h = img.size
    min_width = cfg('ocr', 'min_width', 540)
    ocr_scale = min_width / w if w < min_width else 1.0

    tab_min_y = h * ocr_scale * 0.08
    tab_max_y = h * ocr_scale * 0.25

    cat_clean = re.sub(r'[\s·\-]', '', category_name)

    for bbox, text, conf in results:
        y_mid = (bbox[0][1] + bbox[2][1]) / 2
        if not (tab_min_y <= y_mid <= tab_max_y):
            continue
        text_clean = re.sub(r'[\s·\-]', '', text.strip())
        if cat_clean == text_clean or (
            len(cat_clean) >= 2
            and cat_clean in text_clean
            and len(text_clean) <= len(cat_clean) + 2
        ):
            x_mid = (bbox[0][0] + bbox[2][0]) / 2
            pos = _img_to_screen(region, img, x_mid, y_mid, ocr_scale)
            return pos, img, results

    for bbox, text, conf in results:
        y_mid = (bbox[0][1] + bbox[2][1]) / 2
        if not (tab_min_y <= y_mid <= tab_max_y):
            continue
        text_clean = re.sub(r'[\s·\-]', '', text.strip())
        if cat_clean in text_clean and len(text_clean) > len(cat_clean) + 2:
            idx = text_clean.index(cat_clean)
            char_ratio = (idx + len(cat_clean) / 2) / len(text_clean)
            x_left = bbox[0][0]
            x_right = bbox[2][0]
            x_est = x_left + (x_right - x_left) * char_ratio
            pos = _img_to_screen(region, img, x_est, y_mid, ocr_scale)
            return pos, img, results

    icon_min_y = h * ocr_scale * 0.25
    icon_max_y = h * ocr_scale * 0.55
    for bbox, text, conf in results:
        y_mid = (bbox[0][1] + bbox[2][1]) / 2
        if not (icon_min_y <= y_mid <= icon_max_y):
            continue
        text_clean = re.sub(r'[\s·\-]', '', text.strip())
        if cat_clean == text_clean:
            x_mid = (bbox[0][0] + bbox[2][0]) / 2
            pos = _img_to_screen(region, img, x_mid, y_mid, ocr_scale)
            print(f"    메인 페이지 아이콘으로 '{category_name}' 감지")
            return pos, img, results

    return None, img, results


def switch_category(region, category_name, engine_name=None, reader=None):
    """배민 앱 카테고리 탭을 OCR로 찾아서 클릭."""
    import pyautogui

    for tab in KNOWN_TABS:
        tab_norm = re.sub(r'[\s·\-]', '', tab)
        cat_norm = re.sub(r'[\s·\-]', '', category_name)
        if cat_norm == tab_norm or cat_norm in tab_norm or tab_norm in cat_norm:
            category_name = tab
            break

    if engine_name is None or reader is None:
        return _switch_category_legacy(region, category_name)

    def _try_click_tab():
        pos, img, results = _find_category_tab(region, category_name, engine_name, reader)
        if pos:
            click_x = pos[0] + random.randint(-5, 5)
            click_y = pos[1] + random.randint(-3, 3)
            _bring_to_front()
            pyautogui.moveTo(click_x, click_y,
                             duration=random.uniform(0.2, 0.4))
            time.sleep(random.uniform(0.1, 0.3))
            pyautogui.click()
            _category_pause()
            time.sleep(random.uniform(1.0, 2.0))
            print(f"  '{category_name}' 카테고리로 전환 완료")
            return True
        if _dismiss_popup(region, results, img):
            print(f"  팝업 닫은 후 재시도...")
            time.sleep(random.uniform(0.5, 1.0))
            pos2, _, _ = _find_category_tab(region, category_name, engine_name, reader)
            if pos2:
                _bring_to_front()
                pyautogui.moveTo(pos2[0] + random.randint(-5, 5),
                                 pos2[1] + random.randint(-3, 3),
                                 duration=random.uniform(0.2, 0.4))
                time.sleep(random.uniform(0.1, 0.3))
                pyautogui.click()
                _category_pause()
                time.sleep(random.uniform(1.0, 2.0))
                print(f"  '{category_name}' 카테고리로 전환 완료")
                return True
        return False

    # 1단계: 현재 화면에서 바로 시도
    print(f"  카테고리 '{category_name}' 탭 찾는 중...")
    if _try_click_tab():
        return True

    # 2단계: 위로 점진 스크롤하면서 탭 찾기 (pull-to-refresh 방지)
    for scroll_batch in range(4):
        print(f"  화면에 탭 안 보임, 위로 스크롤 중... ({scroll_batch + 1}/4)")
        _scroll_up(region, count=3)
        if _try_click_tab():
            return True

    # 3단계: 탭바를 오른쪽으로 리셋 후 왼쪽 스와이프하며 찾기
    print(f"  탭바 가로 스크롤 리셋 중...")
    for _ in range(3):
        _swipe_tab_bar(region, 'right')
    time.sleep(random.uniform(0.3, 0.5))

    if _try_click_tab():
        return True

    for swipe in range(4):
        print(f"  탭바 왼쪽 스와이프... ({swipe + 1}/4)")
        _swipe_tab_bar(region, 'left')
        if _try_click_tab():
            return True

    print(f"  카테고리 '{category_name}'을 찾지 못했습니다.")
    print(f"  알려진 카테고리: {', '.join(KNOWN_TABS)}")
    return False


def _switch_category_legacy(region, category_name):
    import pyautogui

    tab_y = region['top'] + int(region['height'] * 0.17)
    tab_x_start = region['left'] + 20
    tab_x_end = region['left'] + region['width'] - 20
    tab_width = (tab_x_end - tab_x_start) // 8

    idx = None
    for i, name in enumerate(KNOWN_TABS):
        if category_name in name or name in category_name:
            idx = i
            break

    if idx is None:
        print(f"  카테고리 '{category_name}'을 목록에서 찾을 수 없습니다.")
        return False

    if idx >= 6:
        for _ in range(3):
            swipe_x_start = tab_x_end - 30
            swipe_x_end = tab_x_start + 30
            pyautogui.moveTo(swipe_x_start, tab_y,
                             duration=random.uniform(0.15, 0.3))
            time.sleep(random.uniform(0.1, 0.2))
            pyautogui.mouseDown()
            pyautogui.moveTo(swipe_x_end, tab_y + random.randint(-3, 3),
                             duration=random.uniform(0.3, 0.5))
            pyautogui.mouseUp()
            time.sleep(random.uniform(0.5, 1.0))
        adjusted_idx = idx - 4
    else:
        adjusted_idx = idx

    click_x = tab_x_start + tab_width * adjusted_idx + tab_width // 2
    click_x += random.randint(-8, 8)
    click_y = tab_y + random.randint(-5, 5)

    pyautogui.moveTo(click_x, click_y, duration=random.uniform(0.2, 0.5))
    time.sleep(random.uniform(0.1, 0.3))
    pyautogui.click()

    _category_pause()
    time.sleep(random.uniform(1.0, 2.0))

    print(f"  '{category_name}' 카테고리로 전환 완료")
    return True


def swipe_down(region):
    import pyautogui
    _bring_to_front()
    x_jitter = random.randint(-20, 20)
    cx = region['left'] + region['width'] // 2 + x_jitter

    start_ratio = random.uniform(0.65, 0.75)
    end_ratio = random.uniform(0.25, 0.35)
    y_start = region['top'] + int(region['height'] * start_ratio)
    y_end = region['top'] + int(region['height'] * end_ratio)

    pyautogui.moveTo(cx, y_start, duration=random.uniform(0.1, 0.3))
    time.sleep(random.uniform(0.05, 0.2))
    pyautogui.mouseDown()
    time.sleep(random.uniform(0.05, 0.15))
    pyautogui.moveTo(cx + random.randint(-5, 5), y_end,
                     duration=random.uniform(0.3, 0.7))
    time.sleep(random.uniform(0.05, 0.15))
    pyautogui.mouseUp()
    human_pause()


# --------------- 중복 제거 ---------------

def _core(name):
    return re.sub(r'[^가-힣a-zA-Z0-9]', '', name)


def _similar(a, b):
    ca, cb = _core(a), _core(b)
    if not ca or not cb:
        return False
    short, long = sorted([ca, cb], key=len)
    if long.startswith(short):
        return True
    if len(short) <= 3:
        hits = sum(1 for c in short if c in long)
        return hits >= len(short) - 1 and len(long) <= len(short) + 1
    hits = sum(1 for c in short if c in long)
    return hits / len(short) >= 0.88


def dedup(businesses):
    out = []
    for b in businesses:
        dup = False
        for existing in out:
            same_score = (b.get('rating') == existing.get('rating')
                          and b.get('reviews') == existing.get('reviews')
                          and b.get('reviews'))
            if _similar(b['name'], existing['name']) or same_score:
                if existing['rating'] is None and b['rating'] is not None:
                    existing['rating'] = b['rating']
                if b['reviews'] and existing['reviews']:
                    if b['reviews'] > existing['reviews']:
                        existing['name'] = b['name']
                        existing['reviews'] = b['reviews']
                        if b['rating'] is not None:
                            existing['rating'] = b['rating']
                dup = True
                break
        if not dup:
            out.append(b)
    return out


# --------------- 저장 ---------------

def save_excel(businesses, filepath):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "배민 업체 목록"

    headers = ["번호", "업체명", "평점", "리뷰 수", "카테고리", "수집 시각"]
    blue = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    white_bold = Font(bold=True, size=11, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = white_bold
        c.fill = blue
        c.alignment = Alignment(horizontal='center')

    for i, biz in enumerate(businesses, 1):
        ws.cell(row=i + 1, column=1, value=i)
        ws.cell(row=i + 1, column=2, value=biz['name'])
        ws.cell(row=i + 1, column=3, value=biz['rating'] if biz['rating'] is not None else '누락')
        ws.cell(row=i + 1, column=4, value=biz['reviews'])
        ws.cell(row=i + 1, column=5, value=biz.get('category', ''))
        ws.cell(row=i + 1, column=6, value=biz['timestamp'])

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 20

    wb.save(str(filepath))


# --------------- 실행 ---------------

def run_test(engine_name, reader, region):
    print("\n[테스트 모드] 화면 한 번 캡처해서 결과를 보여줍니다.\n")

    img = capture(region)
    img_path = OUTPUT_DIR / "test_capture.png"
    img.save(str(img_path))
    print(f"  캡처 저장: {img_path}")

    print("  분석 중...")
    results = read_image(engine_name, reader, img)

    print(f"\n  찾은 텍스트 {len(results)}개:")
    print("  " + "-" * 50)
    for bbox, text, conf in results:
        y = int((bbox[0][1] + bbox[2][1]) / 2)
        print(f"  y={y:4d}  ({conf:.0%})  {text}")
    print("  " + "-" * 50)

    businesses, null_pos, low_conf_pos = parse_businesses(results, img.size[1], return_positions=True)
    print(f"\n  인식된 업체 {len(businesses)}개:")
    for i, b in enumerate(businesses, 1):
        r = f"  {b['rating']}" if b['rating'] else "  (평점누락)"
        rv = f"  리뷰 {b['reviews']}" if b['reviews'] else ""
        low = "  [저신뢰]" if b['name'] in low_conf_pos else ""
        null = "  [평점누락->재시도대상]" if b['name'] in null_pos else ""
        print(f"  {i}. {b['name']}{r}{rv}{low}{null}")


def _crop_and_reocr(reader, img, y_pos, pad=20, scale=3):
    """지정 Y좌표 주변을 크롭+확대하여 재OCR. 결과 리스트 반환."""
    import numpy as np
    w, h = img.size
    top = max(0, int(y_pos) - pad)
    bot = min(h, int(y_pos) + pad)
    crop = img.crop((0, top, w, bot))
    crop = crop.resize((w * scale, (bot - top) * scale), resample=3)
    arr = np.array(crop)
    time.sleep(random.uniform(0.2, 0.5))
    return reader.readtext(arr, mag_ratio=1.5)


def _retry_rating(reader, img, y_pos):
    results = _crop_and_reocr(reader, img, y_pos)
    rating_re = re.compile(r'(\d[.,]\d)')
    for _, text, _ in results:
        m = rating_re.search(text)
        if m:
            val = float(m.group(1).replace(',', '.'))
            if 1.0 <= val <= 5.0:
                return val
    return None


def _has_korean(text):
    return bool(re.search(r'[가-힣]', text))


def _retry_name(reader, img, y_pos):
    results = _crop_and_reocr(reader, img, y_pos, pad=25, scale=3)
    biz_re = re.compile(
        r'(.+?)\s*[★☆\*]?\s*(\d[.,]\d)\s*[\(\[（]\s*([\d,.]+)\s*[\)\]）]'
    )
    for _, text, conf in results:
        if conf < 0.5:
            continue
        m = biz_re.search(text)
        if m:
            name = m.group(1).strip()
            name = re.sub(r'\s*광고\s*$', '', name).strip()
            if len(name) >= 2 and _has_korean(name):
                return name
    long_texts = [(t, c) for _, t, c in results
                  if len(t.strip()) >= 2 and c >= 0.6 and _has_korean(t.strip())]
    if long_texts:
        long_texts.sort(key=lambda x: -x[1])
        return long_texts[0][0].strip()
    return None


_popup_close_re = re.compile(
    r'오늘\s*하루\s*안\s*보기|오늘\s*그만\s*보기|다시\s*보지\s*않기|'
    r'다음에\s*보기|광고\s*닫기'
)
_popup_close_exact_re = re.compile(
    r'^닫기$|^[Xx×]$|^건너뛰기$|^나중에$|^SKIP$'
)
_popup_indicator_re = re.compile(
    r'쿠폰\s*받으러|핫딜|쿠폰\s*도착|이벤트\s*참여|출석\s*체크|'
    r'지금\s*받기|무료\s*받기|앱\s*평가'
)


def _dismiss_popup(region, ocr_results, img):
    """팝업/광고 감지 시 닫기 버튼을 찾아 클릭. 닫았으면 True."""
    if not ocr_results:
        return False

    w, h = img.size
    min_width = cfg('ocr', 'min_width', 540)
    ocr_scale = min_width / w if w < min_width else 1.0

    best = None
    best_priority = 999
    has_popup_indicator = False

    for bbox, text, conf in ocr_results:
        text_clean = text.strip()
        y_mid = (bbox[0][1] + bbox[2][1]) / 2
        x_mid = (bbox[0][0] + bbox[2][0]) / 2

        if _popup_indicator_re.search(text_clean):
            has_popup_indicator = True

        if _popup_close_re.search(text_clean):
            p = 0
        elif _popup_close_exact_re.match(text_clean):
            p = 1
        else:
            continue

        if best is None or p < best_priority or (p == best_priority and y_mid > best[1]):
            best = (x_mid, y_mid, text_clean)
            best_priority = p

    if best:
        x, y, text = best
        sx, sy = _img_to_screen(region, img, x, y, ocr_scale)
        screen_x = sx + random.randint(-3, 3)
        screen_y = sy + random.randint(-2, 2)
        import pyautogui
        _bring_to_front()
        pyautogui.moveTo(screen_x, screen_y, duration=random.uniform(0.2, 0.4))
        time.sleep(random.uniform(0.1, 0.2))
        pyautogui.click()
        time.sleep(random.uniform(1.0, 2.0))
        print(f"  팝업 닫기: '{text}'")
        return True

    if has_popup_indicator:
        scaled_h = h * ocr_scale
        x_close = w * ocr_scale * 0.85
        y_close = scaled_h * 0.25
        sx, sy = _img_to_screen(region, img, x_close, y_close, ocr_scale)
        import pyautogui
        _bring_to_front()
        pyautogui.moveTo(sx + random.randint(-3, 3), sy + random.randint(-2, 2),
                         duration=random.uniform(0.2, 0.4))
        time.sleep(random.uniform(0.1, 0.2))
        pyautogui.click()
        time.sleep(random.uniform(1.5, 2.5))
        print(f"  팝업 감지 -> X 버튼 위치 클릭 시도")
        return True

    return False


def run_collect(engine_name, reader, region, target, category=None):
    max_no_new = cfg('collect', 'max_no_new', 3)
    cat_label = f" [{category}]" if category else ""
    print(f"\n{target}개 업체 수집을 시작합니다.{cat_label}")
    print("마우스를 화면 왼쪽 위 구석으로 옮기면 긴급 중단됩니다.\n")

    all_biz = []
    no_new = 0

    for step in range(1, 500):
        time.sleep(random.uniform(0.5, 1.5))
        print(f"[{step}] 화면 캡처 + 분석 중...")
        img = capture(region)
        img.save(str(OUTPUT_DIR / f"step_{step:03d}.png"))

        results = read_image(engine_name, reader, img)

        for _popup_try in range(3):
            if not _dismiss_popup(region, results, img):
                break
            img = capture(region)
            img.save(str(OUTPUT_DIR / f"step_{step:03d}.png"))
            results = read_image(engine_name, reader, img)

        found, null_pos, low_conf_pos = parse_businesses(results, img.size[1], return_positions=True)

        if engine_name == 'easyocr':
            if null_pos:
                for name, y_pos in null_pos.items():
                    time.sleep(random.uniform(0.1, 0.4))
                    rating = _retry_rating(reader, img, y_pos)
                    if rating is not None:
                        for b in found:
                            if b['name'] == name and b['rating'] is None:
                                b['rating'] = rating
                                print(f"  평점 재시도 성공: {name} -> {rating}")
                                break

            # 이름 보정은 crop 영역에서 인접 업체 텍스트를 읽어올 위험이 있어 비활성화
            # if low_conf_pos:
            #     for old_name, y_pos in low_conf_pos.items():
            #         ...pass

        if category:
            for b in found:
                b['category'] = category

        print(f"  이번 화면: {len(found)}개 업체")

        before = len(all_biz)
        all_biz.extend(found)
        all_biz = dedup(all_biz)
        added = len(all_biz) - before
        print(f"  새로 추가: {added}개  (누적 {len(all_biz)}/{target})")

        if added == 0:
            no_new += 1
            if no_new >= max_no_new:
                print(f"\n  {max_no_new}번 연속 새 업체 없음 -> 수집 종료")
                break
        else:
            no_new = 0

        if len(all_biz) >= target:
            print("\n  목표 달성!")
            break

        print("  스크롤 중...")
        swipe_down(region)

    all_biz = all_biz[:target]
    return all_biz


def save_results(all_biz, label=""):
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    prefix = f"배민_{label}_" if label else "배민_"
    xlsx = OUTPUT_DIR / f"{prefix}{ts}.xlsx"
    save_excel(all_biz, xlsx)

    json_path = OUTPUT_DIR / f"{prefix}{ts}.json"
    with open(str(json_path), 'w', encoding='utf-8') as f:
        json.dump(all_biz, f, ensure_ascii=False, indent=2)

    print(f"\n{'=' * 50}")
    print(f"  수집 완료: {len(all_biz)}개 업체")
    print(f"  엑셀: {xlsx}")
    print(f"  JSON: {json_path}")
    print(f"{'=' * 50}")


def _parse_time(s):
    """'HH:MM' 형태 문자열을 (hour, minute) 튜플로 파싱."""
    parts = s.strip().split(':')
    if len(parts) != 2:
        return None
    try:
        h, m = int(parts[0]), int(parts[1])
        if 0 <= h <= 23 and 0 <= m <= 59:
            return (h, m)
    except ValueError:
        pass
    return None


def _wait_until(hour, minute):
    """지정 시각까지 대기. 이미 지났으면 다음날."""
    now = datetime.now()
    target = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if target <= now:
        target = target.replace(day=target.day + 1)
    delta = (target - now).total_seconds()
    print(f"\n  다음 실행 예정: {target.strftime('%Y-%m-%d %H:%M')}")
    print(f"  대기 중... ({delta / 60:.0f}분)")
    time.sleep(delta)


def _run_once(engine_name, reader, region, count, categories):
    """한 번의 수집 사이클 실행."""
    region = find_bluestacks()
    if not region:
        print("BlueStacks 창을 찾을 수 없습니다. 이번 회차 건너뜁니다.")
        return

    if categories:
        all_results = []
        for ci, cat in enumerate(categories):
            print(f"\n{'=' * 50}")
            print(f"  카테고리 {ci + 1}/{len(categories)}: {cat}")
            print(f"{'=' * 50}")

            if not switch_category(region, cat, engine_name, reader):
                print(f"  '{cat}' 전환 실패, 건너뜁니다.")
                continue

            results = run_collect(engine_name, reader, region, count, category=cat)
            all_results.extend(results)
            print(f"  '{cat}' 수집 완료: {len(results)}개")

            if ci < len(categories) - 1:
                between_wait = random.uniform(5.0, 12.0)
                print(f"\n  다음 카테고리 전환 전 잠시 쉬는 중... {between_wait:.0f}초")
                time.sleep(between_wait)

        save_results(all_results, "멀티")
    else:
        results = run_collect(engine_name, reader, region, count)
        save_results(results)


def main():
    test_mode = '--test' in sys.argv

    count = cfg('collect', 'default_count', 30)
    for i, arg in enumerate(sys.argv):
        if arg == '--count' and i + 1 < len(sys.argv):
            try:
                count = int(sys.argv[i + 1])
            except ValueError:
                pass

    categories = []
    for i, arg in enumerate(sys.argv):
        if arg == '--categories':
            for j in range(i + 1, len(sys.argv)):
                if sys.argv[j].startswith('--'):
                    break
                categories.append(sys.argv[j])

    schedule_time = None
    repeat_minutes = 0
    for i, arg in enumerate(sys.argv):
        if arg == '--schedule' and i + 1 < len(sys.argv):
            schedule_time = _parse_time(sys.argv[i + 1])
            if not schedule_time:
                print(f"시간 형식 오류: '{sys.argv[i + 1]}' (HH:MM 형태로 입력)")
                return
        if arg == '--repeat' and i + 1 < len(sys.argv):
            try:
                repeat_minutes = int(sys.argv[i + 1])
            except ValueError:
                print(f"반복 간격 오류: '{sys.argv[i + 1]}' (분 단위 숫자)")
                return

    if not test_mode and '--count' not in sys.argv and not categories \
       and not schedule_time and repeat_minutes == 0:
        try:
            val = input(f"수집할 업체 수 (기본 {count}): ").strip()
            if val:
                count = int(val)
        except (ValueError, EOFError):
            pass

    print("\nBlueStacks 창을 찾는 중...")
    region = find_bluestacks()
    if not region:
        print("BlueStacks 창을 찾을 수 없습니다.")
        print("BlueStacks를 실행하고 배민 업체 목록을 띄운 뒤 다시 시도하세요.")
        return

    print(f"BlueStacks 발견: {region['title']}")
    print(f"  위치 ({region['left']}, {region['top']})  "
          f"크기 {region['width']}x{region['height']}")

    print("\n엔진 로딩 중...")
    engine_name, reader = create_reader()
    print("준비 완료")

    if test_mode:
        run_test(engine_name, reader, region)
        return

    if schedule_time:
        h, m = schedule_time
        print(f"\n[예약 모드] 매일 {h:02d}:{m:02d}에 수집합니다. (Ctrl+C로 종료)")
        round_num = 0
        while True:
            _wait_until(h, m)
            round_num += 1
            print(f"\n{'#' * 50}")
            print(f"  예약 실행 #{round_num}  {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            print(f"{'#' * 50}")
            startup_wait = random.uniform(10, 60)
            print(f"  시작 전 잠시 대기... {startup_wait:.0f}초")
            time.sleep(startup_wait)
            try:
                _run_once(engine_name, reader, region, count, categories)
            except Exception as e:
                print(f"  오류 발생: {e}")
                print("  다음 예약까지 대기합니다.")
        return

    if repeat_minutes > 0:
        print(f"\n[반복 모드] {repeat_minutes}분 간격으로 수집합니다. (Ctrl+C로 종료)")
        round_num = 0
        while True:
            round_num += 1
            print(f"\n{'#' * 50}")
            print(f"  반복 #{round_num}  {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            print(f"{'#' * 50}")
            try:
                _run_once(engine_name, reader, region, count, categories)
            except Exception as e:
                print(f"  오류 발생: {e}")

            jitter = random.uniform(-60, 120)
            wait_sec = repeat_minutes * 60 + jitter
            wait_sec = max(60, wait_sec)
            print(f"\n  다음 수집까지 {wait_sec / 60:.1f}분 대기...")
            time.sleep(wait_sec)
        return

    print(f"\n3초 뒤 시작합니다. 배민 업체 목록이 보이는지 확인하세요...")
    time.sleep(3)

    if categories:
        print(f"\n{len(categories)}개 카테고리 연속 수집: {', '.join(categories)}")
        print(f"각 카테고리별 {count}개씩 수집합니다.")

        all_results = []
        for ci, cat in enumerate(categories):
            print(f"\n{'=' * 50}")
            print(f"  카테고리 {ci + 1}/{len(categories)}: {cat}")
            print(f"{'=' * 50}")

            if not switch_category(region, cat, engine_name, reader):
                print(f"  '{cat}' 전환 실패, 건너뜁니다.")
                continue

            results = run_collect(engine_name, reader, region, count, category=cat)
            all_results.extend(results)
            print(f"  '{cat}' 수집 완료: {len(results)}개")

            if ci < len(categories) - 1:
                between_wait = random.uniform(5.0, 12.0)
                print(f"\n  다음 카테고리 전환 전 잠시 쉬는 중... {between_wait:.0f}초")
                time.sleep(between_wait)

        save_results(all_results, "멀티")
    else:
        results = run_collect(engine_name, reader, region, count)
        save_results(results)


if __name__ == '__main__':
    main()
